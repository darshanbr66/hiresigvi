"""Deterministic, rule-based extraction of assessment questions from an
uploaded PDF, DOCX, XLSX, or CSV file — shared by the assessment form's
"Import Questions" (docs/assessment.md § 9) and the question bank's import
(docs/assessment.md § 10a). No AI/LLM involved — this never guesses at a
correct answer it can't find explicitly in the document (CLAUDE.md § 5: no
fake implementations); a question whose correct answer can't be determined
is skipped and surfaced as a warning rather than silently invented.

Callers get back plain data (not persisted) so the recruiter can preview,
edit, reorder, and select questions before anything is saved — this module
never touches the database.

Supported conventions:
  - CSV / XLSX: a header row with a `question`/`prompt` column, one or more
    `option`/`option_N` (or `a`, `b`, `c`, `d`, ...) columns, and a
    `correct`/`answer` column identifying the right option(s) by letter,
    1-based index, or exact text match. Optional columns: `type` (single,
    multiple, true_false, short_answer, descriptive), `points`/`marks`,
    `difficulty`, `subject`, `topic`, `tags` (`;`-separated) and
    `explanation`. For `short_answer`, `correct` lists the accepted answers
    separated by `|`; a `descriptive` row needs neither options nor answer.
  - PDF / DOCX: numbered questions ("1. ...") followed by lettered options
    ("A) ...", "B) ..."), with the correct option(s) marked either by a
    trailing "*" on the option line or an explicit "Answer: B" /
    "Correct: B, C" line.

Uploads are untrusted: size, row/page counts and (for DOCX/XLSX) the
decompressed archive size are bounded, the file's content must match its
extension, macros are never run (python-docx/openpyxl don't execute them),
and error messages never echo library exception text.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field

from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
from pypdf import PdfReader

from app.core.exceptions import AppError

#: Largest accepted upload.
MAX_IMPORT_BYTES = 5 * 1024 * 1024
#: Most questions (rows / numbered blocks) read from one file.
MAX_IMPORT_QUESTIONS = 500
#: Most PDF pages read.
MAX_PDF_PAGES = 100
#: Ceiling on the total uncompressed size of a DOCX/XLSX archive (zip-bomb guard).
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 50 * 1024 * 1024

SUPPORTED_EXTENSIONS = ("pdf", "docx", "xlsx", "csv")


@dataclass
class ParsedOption:
    label: str
    is_correct: bool = False


@dataclass
class ParsedQuestion:
    prompt: str
    type: str  # "MCQ_SINGLE" | "MCQ_MULTI" | "TRUE_FALSE" | "SHORT_ANSWER" | "DESCRIPTIVE"
    points: int = 1
    options: list[ParsedOption] = field(default_factory=list)
    #: Where it came from ("Row 4", "Question 2") — for row-specific messages.
    source: str = ""
    accepted_answers: list[str] = field(default_factory=list)
    # Optional organising metadata (tabular files only).
    difficulty: str | None = None
    subject: str | None = None
    topic: str | None = None
    tags: list[str] = field(default_factory=list)
    explanation: str | None = None


@dataclass
class QuestionImportResult:
    questions: list[ParsedQuestion]
    warnings: list[str]


class QuestionImportError(AppError):
    code = "question_import_failed"


def parse_questions(*, content: bytes, filename: str) -> QuestionImportResult:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix not in SUPPORTED_EXTENSIONS:
        raise QuestionImportError(
            f"Unsupported file type: .{suffix or 'unknown'}. Upload a PDF, DOCX, XLSX, or CSV file."
        )
    if not content:
        raise QuestionImportError("The file is empty.")
    if len(content) > MAX_IMPORT_BYTES:
        raise QuestionImportError(
            f"The file is larger than {MAX_IMPORT_BYTES // (1024 * 1024)} MB. "
            "Split it into smaller files."
        )
    _check_signature(content, suffix)

    if suffix == "csv":
        return _parse_tabular(_read_csv_rows(content))
    if suffix == "xlsx":
        return _parse_tabular(_read_xlsx_rows(content))
    if suffix == "docx":
        return _parse_freeform_text(_read_docx_text(content))
    return _parse_freeform_text(_read_pdf_text(content))


# --- Content checks -------------------------------------------------------

_ARCHIVE_MARKERS = {"docx": "word/document.xml", "xlsx": "xl/workbook.xml"}


def _mismatch(suffix: str) -> QuestionImportError:
    return QuestionImportError(
        f"This file's contents don't match a .{suffix} file. Save it again as .{suffix} and retry."
    )


def _check_signature(content: bytes, suffix: str) -> None:
    """The extension alone is never trusted."""
    if suffix == "pdf":
        if not content.lstrip()[:5] == b"%PDF-":
            raise _mismatch(suffix)
        return
    if suffix == "csv":
        if b"\x00" in content[:8192] or content[:4] in (b"PK\x03\x04", b"%PDF"):
            raise _mismatch(suffix)
        return
    # DOCX / XLSX are zip archives.
    if content[:4] != b"PK\x03\x04":
        raise _mismatch(suffix)
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = set(archive.namelist())
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise _mismatch(suffix) from exc
    if _ARCHIVE_MARKERS[suffix] not in names:
        raise _mismatch(suffix)
    if total > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
        raise QuestionImportError("This file expands to more data than an import can handle.")


# --- File readers ---------------------------------------------------------


def _read_csv_rows(content: bytes) -> list[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise QuestionImportError(
            "Could not read this CSV file as text. Save it as 'CSV UTF-8' and try again."
        ) from exc
    try:
        reader = csv.reader(io.StringIO(text))
        return [row for row in reader if any(cell.strip() for cell in row)]
    except csv.Error as exc:
        raise QuestionImportError(
            "This CSV file is malformed (check for unbalanced quotes)."
        ) from exc


def _read_xlsx_rows(content: bytes) -> list[list[str]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(["" if cell is None else str(cell) for cell in row])
                if len(rows) > MAX_IMPORT_QUESTIONS + 1:
                    break
        finally:
            workbook.close()
    except QuestionImportError:
        raise
    except Exception as exc:
        raise QuestionImportError(
            "Could not read this XLSX file. Check that it opens in Excel."
        ) from exc
    return [row for row in rows if any(cell.strip() for cell in row)]


def _read_docx_text(content: bytes) -> str:
    try:
        document = Document(io.BytesIO(content))
        numbering = _DocxNumbering(document)
        lines = [numbering.marker(paragraph) + paragraph.text for paragraph in document.paragraphs]
    except Exception as exc:
        raise QuestionImportError(
            "Could not read this DOCX file. Check that it opens in Word."
        ) from exc
    text = "\n".join(lines).strip()
    if not text:
        raise QuestionImportError("No extractable text found in this DOCX file.")
    return text


class _DocxNumbering:
    """Word's automatic list numbering isn't part of `paragraph.text` — a
    question typed as "1." in Word is usually a "List Number" paragraph whose
    "1." only exists in numbering.xml. This rebuilds the visible marker
    ("3. " for decimal lists, "B) " for lettered ones) so auto-numbered
    documents parse exactly like typed ones. Other formats (bullets, roman
    numerals) get no marker."""

    _LETTER_FORMATS = ("lowerLetter", "upperLetter")

    def __init__(self, document: object) -> None:
        self._formats: dict[tuple[str, int], tuple[str, int]] = {}
        self._counters: dict[str, dict[int, int]] = {}
        try:
            root = document.part.numbering_part.element  # type: ignore[attr-defined]
        except (KeyError, NotImplementedError):
            return
        abstract_levels: dict[str, dict[int, tuple[str, int]]] = {}
        for abstract in root.findall(qn("w:abstractNum")):
            levels: dict[int, tuple[str, int]] = {}
            for lvl in abstract.findall(qn("w:lvl")):
                fmt = lvl.find(qn("w:numFmt"))
                start = lvl.find(qn("w:start"))
                levels[int(lvl.get(qn("w:ilvl"), "0"))] = (
                    fmt.get(qn("w:val"), "decimal") if fmt is not None else "decimal",
                    int(start.get(qn("w:val"), "1")) if start is not None else 1,
                )
            abstract_levels[abstract.get(qn("w:abstractNumId"), "")] = levels
        for num in root.findall(qn("w:num")):
            ref = num.find(qn("w:abstractNumId"))
            if ref is None:
                continue
            for ilvl, spec in abstract_levels.get(ref.get(qn("w:val"), ""), {}).items():
                self._formats[(num.get(qn("w:numId"), ""), ilvl)] = spec

    @staticmethod
    def _num_pr(paragraph: object) -> tuple[str, int] | None:
        """(numId, ilvl) from the paragraph, else from its style chain."""
        sources = [paragraph._p.pPr]  # type: ignore[attr-defined]
        style = paragraph.style  # type: ignore[attr-defined]
        while style is not None:
            sources.append(style.element.pPr)
            style = style.base_style
        num_id: str | None = None
        ilvl: int | None = None
        for p_pr in sources:
            num_pr = p_pr.numPr if p_pr is not None else None
            if num_pr is None:
                continue
            if num_id is None and num_pr.numId is not None:
                num_id = str(num_pr.numId.val)
            if ilvl is None and num_pr.ilvl is not None:
                ilvl = int(num_pr.ilvl.val)
        if num_id is None or num_id == "0":
            return None
        return num_id, ilvl or 0

    def marker(self, paragraph: object) -> str:
        num_pr = self._num_pr(paragraph)
        if num_pr is None or num_pr not in self._formats:
            return ""
        num_id, ilvl = num_pr
        fmt, start = self._formats[num_pr]
        counters = self._counters.setdefault(num_id, {})
        value = counters.get(ilvl, start - 1) + 1
        counters[ilvl] = value
        # A new item at this level restarts every deeper level (options
        # under question 2 are lettered from A again).
        for deeper in [level for level in counters if level > ilvl]:
            del counters[deeper]
        if fmt == "decimal":
            return f"{value}. "
        if fmt in self._LETTER_FORMATS and 1 <= value <= 26:
            return f"{chr(ord('A') + value - 1)}) "
        return ""


def _read_pdf_text(content: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise QuestionImportError(
                "This PDF is password-protected. Remove the password and try again."
            )
        if len(reader.pages) > MAX_PDF_PAGES:
            raise QuestionImportError(
                f"This PDF has more than {MAX_PDF_PAGES} pages. Split it into smaller files."
            )
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except QuestionImportError:
        raise
    except Exception as exc:
        raise QuestionImportError(
            "Could not read this PDF file. Check that it opens normally."
        ) from exc
    text = text.strip()
    if not text:
        raise QuestionImportError(
            "No extractable text found in this PDF (it may be a scanned image)."
        )
    return text


# --- Tabular (CSV/XLSX) parsing -------------------------------------------

_OPTION_COLUMN_RE = re.compile(r"^option[_ ]?(\d+)$")
_LETTER_OPTION_COLUMNS = ["a", "b", "c", "d", "e", "f"]
_DIFFICULTIES = {"easy": "EASY", "medium": "MEDIUM", "hard": "HARD"}


def _normalize_header(raw: str) -> str:
    return re.sub(r"[\s\-]+", "_", raw.strip().lower())


_TYPE_ALIASES = {
    "single": "MCQ_SINGLE",
    "single_choice": "MCQ_SINGLE",
    "mcq": "MCQ_SINGLE",
    "mcq_single": "MCQ_SINGLE",
    "multi": "MCQ_MULTI",
    "multiple": "MCQ_MULTI",
    "multiple_choice": "MCQ_MULTI",
    "multi_select": "MCQ_MULTI",
    "mcq_multi": "MCQ_MULTI",
    "true_false": "TRUE_FALSE",
    "truefalse": "TRUE_FALSE",
    "tf": "TRUE_FALSE",
    "boolean": "TRUE_FALSE",
    "short": "SHORT_ANSWER",
    "short_answer": "SHORT_ANSWER",
    "descriptive": "DESCRIPTIVE",
    "essay": "DESCRIPTIVE",
    "long_answer": "DESCRIPTIVE",
}


def _declared_type(raw: str) -> str | None:
    """Maps the optional `type` cell; None = infer from the answer key,
    "UNKNOWN" = a value we don't recognise (reported, never guessed)."""
    value = re.sub(r"[\s\-/]+", "_", raw.strip().lower())
    if not value:
        return None
    return _TYPE_ALIASES.get(value, "UNKNOWN")


def _parse_tabular(rows: list[list[str]]) -> QuestionImportResult:
    if not rows:
        raise QuestionImportError("This file has no rows.")

    header = [_normalize_header(cell) for cell in rows[0]]

    def column(*names: str) -> int | None:
        return next((i for i, h in enumerate(header) if h in names), None)

    question_idx = column("question", "prompt")
    if question_idx is None:
        raise QuestionImportError(
            "Expected a 'question' column in the first row. "
            "See the import template for the expected format."
        )

    option_cols: list[tuple[int, int]] = []  # (column index, option number)
    for i, h in enumerate(header):
        match = _OPTION_COLUMN_RE.match(h)
        if match:
            option_cols.append((i, int(match.group(1))))
    if not option_cols:
        for i, h in enumerate(header):
            if h in _LETTER_OPTION_COLUMNS:
                option_cols.append((i, _LETTER_OPTION_COLUMNS.index(h) + 1))
    option_cols.sort(key=lambda pair: pair[1])

    correct_idx = column("correct", "answer", "correct_answer")
    type_idx = column("type", "question_type")
    points_idx = column("points", "marks")
    difficulty_idx = column("difficulty")
    subject_idx = column("subject", "field")
    topic_idx = column("topic", "subtopic")
    tags_idx = column("tags", "tag")
    explanation_idx = column("explanation", "rationale")

    if not option_cols and type_idx is None:
        raise QuestionImportError(
            "Expected option columns such as 'option_1', 'option_2', ... (or 'a', 'b', 'c', ...)."
        )

    data_rows = rows[1:]
    warnings: list[str] = []
    if len(data_rows) > MAX_IMPORT_QUESTIONS:
        warnings.append(
            f"Only the first {MAX_IMPORT_QUESTIONS} rows were read; split larger files."
        )
        data_rows = data_rows[:MAX_IMPORT_QUESTIONS]

    questions: list[ParsedQuestion] = []
    for row_number, row in enumerate(data_rows, start=2):

        def cell(i: int | None, row: list[str] = row) -> str:
            return row[i].strip() if i is not None and i < len(row) and row[i] is not None else ""

        prompt = cell(question_idx)
        if not prompt:
            continue
        source = f"Row {row_number}"

        declared = _declared_type(cell(type_idx))
        if declared == "UNKNOWN":
            warnings.append(
                f'{source}: skipped — unknown type "{cell(type_idx)}" (use single, multiple, '
                "true_false, short_answer or descriptive)."
            )
            continue

        points_raw = cell(points_idx)
        if points_raw and not (points_raw.isdigit() and 1 <= int(points_raw) <= 1000):
            warnings.append(f"{source}: skipped — marks must be a whole number from 1 to 1000.")
            continue
        points = int(points_raw) if points_raw else 1

        difficulty_raw = cell(difficulty_idx).lower()
        if difficulty_raw and difficulty_raw not in _DIFFICULTIES:
            warnings.append(f"{source}: skipped — difficulty must be easy, medium or hard.")
            continue

        metadata = {
            "difficulty": _DIFFICULTIES.get(difficulty_raw),
            "subject": cell(subject_idx) or None,
            "topic": cell(topic_idx) or None,
            "tags": [t.strip() for t in re.split(r"[;|]", cell(tags_idx)) if t.strip()],
            "explanation": cell(explanation_idx) or None,
        }

        if declared in ("SHORT_ANSWER", "DESCRIPTIVE"):
            accepted = (
                [a.strip() for a in cell(correct_idx).split("|") if a.strip()]
                if declared == "SHORT_ANSWER"
                else []
            )
            questions.append(
                ParsedQuestion(
                    prompt=prompt,
                    type=declared,
                    points=points,
                    source=source,
                    accepted_answers=accepted,
                    **metadata,
                )
            )
            continue

        options = [cell(col_idx) for col_idx, _ in option_cols if cell(col_idx)]
        if len(options) < 2:
            warnings.append(f"{source}: skipped — fewer than 2 non-empty options.")
            continue

        raw_correct = cell(correct_idx)
        correct_indices = _resolve_correct_indices(raw=raw_correct, options=options)
        if not correct_indices:
            reason = (
                "no correct answer given"
                if not raw_correct
                else f'the answer "{raw_correct}" doesn\'t match any option'
            )
            warnings.append(f"{source}: skipped — {reason}.")
            continue

        if declared is not None:
            q_type = declared
        else:
            q_type = "MCQ_MULTI" if len(correct_indices) > 1 else "MCQ_SINGLE"

        questions.append(
            ParsedQuestion(
                prompt=prompt,
                type=q_type,
                points=points,
                source=source,
                options=[
                    ParsedOption(label=label, is_correct=(i in correct_indices))
                    for i, label in enumerate(options)
                ],
                **metadata,
            )
        )

    return QuestionImportResult(questions=questions, warnings=warnings)


def _resolve_correct_indices(*, raw: str, options: list[str]) -> set[int]:
    """`raw` may name correct options by 1-based index ("2"), by letter
    ("B"), or by exact text match — comma/semicolon/'and'-separated for
    multiple answers."""
    if not raw:
        return set()

    tokens = [t.strip() for t in re.split(r"[,;/]| and ", raw, flags=re.IGNORECASE) if t.strip()]
    indices: set[int] = set()
    for token in tokens:
        if token.isdigit():
            i = int(token) - 1
            if 0 <= i < len(options):
                indices.add(i)
            continue
        if len(token) == 1 and token.isalpha():
            i = ord(token.upper()) - ord("A")
            if 0 <= i < len(options):
                indices.add(i)
                continue
        for i, label in enumerate(options):
            if label.strip().lower() == token.lower():
                indices.add(i)
                break
    return indices


# --- Free-text (PDF/DOCX) parsing -----------------------------------------

_QUESTION_START_RE = re.compile(r"^\s*(\d+)[.)]\s+(.*\S)\s*$")
_OPTION_LINE_RE = re.compile(r"^\s*\(?([A-Za-z])[.)]\s+(.*\S)\s*$")
_ANSWER_LINE_RE = re.compile(r"^\s*(?:answer|correct)s?\s*[:\-]\s*(.+\S)\s*$", re.IGNORECASE)
_POINTS_SUFFIX_RE = re.compile(r"[\[(]\s*(\d+)\s*(?:pts?|points?|marks?)\s*[\])]", re.IGNORECASE)


def _parse_freeform_text(text: str) -> QuestionImportResult:
    lines = [line.rstrip() for line in text.splitlines()]

    blocks: list[list[str]] = []
    current: list[str] | None = None
    for line in lines:
        if _QUESTION_START_RE.match(line):
            if current is not None:
                blocks.append(current)
            current = [line]
        elif current is not None:
            current.append(line)
    if current is not None:
        blocks.append(current)

    if not blocks:
        raise QuestionImportError(
            "No numbered questions (e.g. '1. What is ...') were found in this document. "
            "See the import template for the expected format."
        )

    warnings: list[str] = []
    if len(blocks) > MAX_IMPORT_QUESTIONS:
        warnings.append(
            f"Only the first {MAX_IMPORT_QUESTIONS} questions were read; split larger files."
        )
        blocks = blocks[:MAX_IMPORT_QUESTIONS]

    questions: list[ParsedQuestion] = []
    for block_number, block in enumerate(blocks, start=1):
        start_match = _QUESTION_START_RE.match(block[0])
        assert start_match is not None
        source = f"Question {block_number}"
        prompt_lines = [start_match.group(2)]

        options: list[ParsedOption] = []
        answer_letters: set[str] | None = None
        in_options = False

        for line in block[1:]:
            if not line.strip():
                continue
            option_match = _OPTION_LINE_RE.match(line)
            answer_match = _ANSWER_LINE_RE.match(line)
            if answer_match:
                answer_letters = {
                    t.strip().upper()
                    for t in re.split(r"[,;/]| and ", answer_match.group(1), flags=re.IGNORECASE)
                    if t.strip()
                }
                continue
            if option_match:
                in_options = True
                label = option_match.group(2).strip()
                is_marked = label.endswith("*")
                if is_marked:
                    label = label[:-1].strip()
                options.append(ParsedOption(label=label, is_correct=is_marked))
                continue
            if not in_options:
                prompt_lines.append(line.strip())

        prompt = " ".join(p for p in prompt_lines if p).strip()
        points_match = _POINTS_SUFFIX_RE.search(prompt)
        points = int(points_match.group(1)) if points_match else 1
        if points_match:
            prompt = _POINTS_SUFFIX_RE.sub("", prompt).strip()

        if len(options) < 2:
            warnings.append(f"{source}: skipped — fewer than 2 options found.")
            continue

        if answer_letters:
            valid_letters = {chr(ord("A") + i) for i in range(len(options))}
            unknown = sorted(answer_letters - valid_letters)
            if unknown:
                warnings.append(
                    f"{source}: skipped — the answer line names option(s) {', '.join(unknown)}, "
                    "which the question doesn't have."
                )
                continue
            for i, option in enumerate(options):
                option.is_correct = chr(ord("A") + i) in answer_letters

        correct_count = sum(1 for o in options if o.is_correct)
        if correct_count == 0:
            warnings.append(
                f"{source}: skipped — no correct answer marked "
                "(mark it with a trailing '*' or an 'Answer: X' line)."
            )
            continue

        questions.append(
            ParsedQuestion(
                prompt=prompt,
                type="MCQ_MULTI" if correct_count > 1 else "MCQ_SINGLE",
                points=points,
                options=options,
                source=source,
            )
        )

    return QuestionImportResult(questions=questions, warnings=warnings)
