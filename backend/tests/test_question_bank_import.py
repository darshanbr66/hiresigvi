"""Question bank file import (docs/assessment.md § 10a): preview saves
nothing, problems are reported per row, likely duplicates are flagged,
confirming is all-or-nothing, subjects/topics are filed as categories, and
permissions and tenant isolation hold."""

import io

from docx import Document
from httpx import AsyncClient
from openpyxl import Workbook

from app.models.user import User
from tests.org_setup import bootstrap_org, staff_headers
from tests.test_question_bank import BANK, _category, _create, mcq

PREVIEW = f"{BANK}/import/preview"
IMPORT = f"{BANK}/import"

CSV = (
    b"question,option_1,option_2,option_3,correct,type,difficulty,subject,topic,marks,tags\n"
    b"Which keyword defines a function in Python?,func,def,lambda,2,single,easy,"
    b"Programming,Python,2,basics\n"
    b'Pick the SQL aggregates,COUNT,JOIN,SUM,"A,C",multiple,medium,Databases,,1,\n'
    b"Unanswerable?,yes,no,,,,,,,\n"
    b"Explain PID control.,,,,,descriptive,hard,Robotics,Control,5,\n"
)


async def _preview(
    client: AsyncClient, ctx: dict, content: bytes, filename: str, headers: dict | None = None
):
    return await client.post(
        PREVIEW,
        files={"file": (filename, content, "application/octet-stream")},
        headers=headers or ctx["headers"],
    )


def _items(preview: dict) -> list[dict]:
    """What the UI sends back for the questions it keeps."""
    return [
        {
            "content": q["content"],
            "difficulty": q["difficulty"],
            "subject": q["subject"],
            "topic": q["topic"],
            "tags": q["tags"],
        }
        for q in preview["questions"]
        if q["content"] is not None and q["issue"] is None
    ]


async def test_template_is_a_real_csv_download(client: AsyncClient, super_admin: User) -> None:
    ctx = await bootstrap_org(client, "qbi-template")
    response = await client.get(f"{BANK}/import/template", headers=ctx["headers"])
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert "attachment" in response.headers["content-disposition"]
    text = response.content.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("question,option_1,option_2,option_3,option_4,correct")
    # The template itself imports cleanly.
    preview = (await _preview(client, ctx, response.content, "template.csv")).json()
    assert preview["warnings"] == []
    assert all(q["issue"] is None for q in preview["questions"])
    assert len(preview["questions"]) == 3


async def test_preview_reports_rows_and_saves_nothing(
    client: AsyncClient, super_admin: User
) -> None:
    ctx = await bootstrap_org(client, "qbi-preview")
    response = await _preview(client, ctx, CSV, "questions.csv")
    assert response.status_code == 200, response.text
    preview = response.json()

    assert [q["source"] for q in preview["questions"]] == ["Row 2", "Row 3", "Row 5"]
    first = preview["questions"][0]
    assert first["content"]["type"] == "MCQ_SINGLE"
    assert first["content"]["points"] == 2
    assert first["category_path"] == "Programming › Python"
    assert first["creates_category"] is True
    assert preview["questions"][2]["content"]["type"] == "DESCRIPTIVE"
    assert preview["warnings"] == ["Row 4: skipped — no correct answer given."]

    listing = await client.get(f"{BANK}/questions", headers=ctx["headers"])
    assert listing.headers["X-Total-Count"] == "0"
    assert (await client.get(f"{BANK}/categories", headers=ctx["headers"])).json() == []


async def test_import_files_questions_into_categories(
    client: AsyncClient, super_admin: User
) -> None:
    ctx = await bootstrap_org(client, "qbi-import")
    programming = await _category(client, ctx, "Programming")
    preview = (await _preview(client, ctx, CSV, "questions.csv")).json()
    # "Programming" exists, so only its "Python" topic is new.
    assert preview["questions"][0]["creates_category"] is True
    assert preview["questions"][0]["category_path"] == "Programming › Python"

    result = await client.post(
        IMPORT, json={"items": _items(preview), "status": "PUBLISHED"}, headers=ctx["headers"]
    )
    assert result.status_code == 201, result.text
    body = result.json()
    assert body["imported"] == 3
    # Python (under Programming), Databases, Robotics, Control (under Robotics).
    assert body["categories_created"] == 4

    categories = {
        c["path"]: c
        for c in (await client.get(f"{BANK}/categories", headers=ctx["headers"])).json()
    }
    assert categories["Programming › Python"]["parent_id"] == programming["id"]
    assert "Robotics › Control" in categories

    listing = (
        await client.get(f"{BANK}/questions?status=PUBLISHED", headers=ctx["headers"])
    ).json()
    assert len(listing) == 3
    item = (
        await client.get(f"{BANK}/questions/{body['item_ids'][0]}", headers=ctx["headers"])
    ).json()
    assert item["difficulty"] == "EASY"
    assert item["tags"] == ["basics"]
    assert [o["is_correct"] for o in item["current_version"]["options"]] == [False, True, False]


async def test_import_is_all_or_nothing(client: AsyncClient, super_admin: User) -> None:
    ctx = await bootstrap_org(client, "qbi-atomic")
    preview = (await _preview(client, ctx, CSV, "questions.csv")).json()
    items = _items(preview)
    # A single-choice question with two correct answers is invalid.
    items.append(
        {
            **items[0],
            "content": {
                **items[0]["content"],
                "options": [{"label": "a", "is_correct": True}, {"label": "b", "is_correct": True}],
            },
        }
    )

    response = await client.post(IMPORT, json={"items": items}, headers=ctx["headers"])
    assert response.status_code == 422
    assert (await client.get(f"{BANK}/questions", headers=ctx["headers"])).headers[
        "X-Total-Count"
    ] == "0"
    assert (await client.get(f"{BANK}/categories", headers=ctx["headers"])).json() == []


async def test_duplicates_are_flagged_for_review(client: AsyncClient, super_admin: User) -> None:
    ctx = await bootstrap_org(client, "qbi-dupes")
    existing = await _create(client, ctx, mcq("Which keyword defines a function in Python?"))
    content = (
        b"question,option_1,option_2,correct\n"
        b"which keyword   defines a function in python?,def,func,1\n"
        b"New question?,a,b,1\n"
        b"New question?,a,b,2\n"
    )
    preview = (await _preview(client, ctx, content, "dupes.csv")).json()
    flags = [(q["duplicate_of_item_id"], q["duplicate_in_file"]) for q in preview["questions"]]
    assert flags == [(existing["id"], False), (None, False), (None, True)]


async def test_docx_xlsx_and_pdf_previews(client: AsyncClient, super_admin: User) -> None:
    from tests.test_question_import import _multiline_pdf

    ctx = await bootstrap_org(client, "qbi-formats")

    document = Document()
    for line in ["1. Capital of France?", "A) Berlin", "B) Paris*"]:
        document.add_paragraph(line)
    docx = io.BytesIO()
    document.save(docx)

    workbook = Workbook()
    workbook.active.append(["question", "option_1", "option_2", "correct", "subject"])
    workbook.active.append(["2 + 2 = ?", "3", "4", "B", "Aptitude"])
    xlsx = io.BytesIO()
    workbook.save(xlsx)

    pdf = _multiline_pdf(["1. Largest planet?", "A) Mars", "B) Jupiter", "Answer: B"])

    for content, name in ((docx.getvalue(), "q.docx"), (xlsx.getvalue(), "q.xlsx"), (pdf, "q.pdf")):
        response = await _preview(client, ctx, content, name)
        assert response.status_code == 200, (name, response.text)
        questions = response.json()["questions"]
        assert len(questions) == 1 and questions[0]["issue"] is None, name
        assert [o["is_correct"] for o in questions[0]["content"]["options"]] == [False, True], name


async def test_bad_files_get_actionable_errors(client: AsyncClient, super_admin: User) -> None:
    ctx = await bootstrap_org(client, "qbi-bad")
    wrong = await _preview(client, ctx, b"question,option_1\n", "q.pdf")
    assert wrong.status_code == 400
    assert "don't match" in wrong.json()["error"]["message"]
    unsupported = await _preview(client, ctx, b"x", "q.txt")
    assert unsupported.status_code == 400
    assert unsupported.json()["error"]["code"] == "question_import_failed"


async def test_import_permissions_and_isolation(client: AsyncClient, super_admin: User) -> None:
    ctx = await bootstrap_org(client, "qbi-perm-a")
    other = await bootstrap_org(client, "qbi-perm-b")
    manager = await staff_headers(client, ctx, "HIRING_MANAGER")

    assert (await _preview(client, ctx, CSV, "q.csv", headers=manager)).status_code == 403
    assert (await client.get(f"{BANK}/import/template", headers=manager)).status_code == 403
    preview = (await _preview(client, ctx, CSV, "q.csv")).json()
    assert (
        await client.post(IMPORT, json={"items": _items(preview)}, headers=manager)
    ).status_code == 403
    assert (
        await client.post(IMPORT, json={"items": _items(preview)}, headers={})
    ).status_code == 401

    # Importing into one organization never touches another's bank or categories.
    assert (
        await client.post(IMPORT, json={"items": _items(preview)}, headers=ctx["headers"])
    ).status_code == 201
    assert (await client.get(f"{BANK}/questions", headers=other["headers"])).headers[
        "X-Total-Count"
    ] == "0"
    assert (await client.get(f"{BANK}/categories", headers=other["headers"])).json() == []
    # ...and duplicate detection only looks inside the caller's organization.
    other_preview = (await _preview(client, other, CSV, "q.csv")).json()
    assert all(q["duplicate_of_item_id"] is None for q in other_preview["questions"])
